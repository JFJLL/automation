import openpyxl
import io
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime, timedelta
from platforms.registry import (
    calculate_match_scores, detect_best_platform,
    find_id_column, find_date_column, PLATFORMS
)

def normalize_date_str(val: Any) -> str:
    s = str(val).strip().replace("/", "-").split(" ")[0]
    digits = "".join(c for c in s if c.isdigit())
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    return s

def normalize_date_compact(val: Any) -> str:
    s = str(val).strip().replace("/", "-").split(" ")[0]
    digits = "".join(c for c in s if c.isdigit())
    if len(digits) == 8:
        return digits
    return s

def parse_excel_sheets(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets_info = []
    for title in wb.sheetnames:
        ws = wb[title]
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            continue
        header_row_idx = 0
        header_vals = []
        for idx, row in enumerate(rows_iter[:10]):
            clean_r = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if len(clean_r) >= 2:
                header_row_idx = idx
                header_vals = [str(c).strip() if c is not None else "" for c in row]
                while header_vals and not header_vals[-1]:
                    header_vals.pop()
                break
        if not header_vals:
            continue
        
        all_data_rows = []
        for r in rows_iter[header_row_idx + 1:]:
            row_dict = {}
            for col_name, val in zip(header_vals, r):
                if col_name:
                    row_dict[col_name] = "" if val is None else str(val).strip()
            if any(row_dict.values()):
                all_data_rows.append(row_dict)
        sample_rows = all_data_rows[:30]
        
        scores = calculate_match_scores(header_vals)
        best_code, best_info = detect_best_platform(header_vals)
        
        sheets_info.append({
            "sheet_title": title,
            "header_row_index": header_row_idx,
            "headers": header_vals,
            "sample_rows": sample_rows,
            "all_data_rows": all_data_rows,
            "detected_platform": best_code,
            "match_scores": scores
        })
    return sheets_info

def analyze_sheet_for_platform(sheet_info: Dict[str, Any], selected_platform: str) -> Dict[str, Any]:
    headers = sheet_info["headers"]
    sample_rows = sheet_info["sample_rows"]
    scores = sheet_info["match_scores"]
    
    selected_score = scores.get(selected_platform, {})
    detected_platform = sheet_info["detected_platform"]
    detected_score = scores.get(detected_platform, {})
    
    needs_confirm = False
    confirm_message = ""
    if detected_platform != selected_platform:
        sel_has_id = selected_score.get("has_id_column", False)
        det_has_id = detected_score.get("has_id_column", False)
        sel_score = selected_score.get("score", 0)
        det_score = detected_score.get("score", 0)
        if (not sel_has_id and det_has_id) or (det_score > sel_score + 5):
            needs_confirm = True
            sel_name = PLATFORMS.get(selected_platform, {}).get("name", selected_platform)
            det_name = PLATFORMS.get(detected_platform, {}).get("name", detected_platform)
            confirm_message = (
                f"您选择了【{sel_name}】（匹配率 {selected_score.get('score', 0)}%），"
                f"但表头特征更符合【{det_name}】（匹配率 {detected_score.get('score', 0)}%）。\n"
                f"若您确认该表格确实来自【{sel_name}】，可点击继续；也可以切换为【{det_name}】。"
            )
    
    id_col = find_id_column(selected_platform, headers)
    date_col = find_date_column(selected_platform, headers)
    
    entity_ids = []
    seen = set()
    dates = []
    all_rows = sheet_info.get("all_data_rows") or sample_rows
    for r in all_rows:
        val = str(r.get(id_col, "")).strip()
        if val and val not in seen and val != id_col:
            seen.add(val)
            entity_ids.append(val)
        d_val = normalize_date_str(r.get(date_col, ""))
        if len(d_val) == 10 and d_val[4] == "-" and d_val[7] == "-":
            dates.append(d_val)
    
    dimension = "默认"
    if selected_platform == "taobao":
        dimension = "内容" if "内容ID" in headers else "任务"
    elif selected_platform == "juguang":
        if "精准定向" in headers:
            dimension = "target"
        elif "关键词" in headers:
            dimension = "keyword"
        elif "投放位置" in headers and "创意ID" not in headers:
            dimension = "account"
        else:
            dimension = "creativity"
            
    platform_vocab = set(PLATFORMS.get(selected_platform, {}).get("vocab", []))
    column_mapping = []
    for h in headers:
        is_key = (h == id_col or h == date_col)
        mapped = (h in platform_vocab)
        column_mapping.append({
            "column_name": h,
            "is_key": is_key,
            "mapped": mapped,
            "status": "ok" if mapped else "unmapped_metric"
        })
        
    return {
        "sheet_title": sheet_info["sheet_title"],
        "headers": headers,
        "id_column": id_col,
        "sample_id_values": entity_ids[:5],
        "date_column": date_col,
        "detected_entity_ids": entity_ids,
        "min_sample_date": min(dates) if dates else None,
        "max_sample_date": max(dates) if dates else None,
        "next_start_date": (datetime.strptime(max(dates), "%Y-%m-%d").date() + timedelta(days=1)).isoformat() if dates else None,
        "uploaded_rows": [[normalize_date_compact(r.get(h, "")) if h == date_col else (r.get(h, "") if r.get(h) is not None else "") for h in headers] for r in (sheet_info.get("all_data_rows") or sample_rows)],
        "dimension": dimension,
        "column_mapping": column_mapping,
        "needs_confirm": needs_confirm,
        "confirm_message": confirm_message,
        "selected_platform": selected_platform,
        "detected_platform": detected_platform
    }
